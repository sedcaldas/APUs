import os
import openpyxl
from decimal import Decimal
import datetime
from sqlalchemy.orm import Session
from backend.app.database import engine, SessionLocal, Base
from backend.app.models import (
    Departamento, Subregion, Municipio, InstitucionEducativa, SedeEducativa,
    Capitulo, Actividad, Insumo, PrecioInsumo, APU, APUComponente,
    AnalisisIndirecto, ParametroPrestacional, PolizaRango,
    ProyectoObra, ProyectoSede, Presupuesto, PresupuestoItem, Usuario
)
from backend.app.services.calculation import recalculate_budget, freeze_budget_snapshot

CALDAS_SUBREGIONES = [
    {
        "codigo": "CS",
        "nombre": "Centrosur",
        "municipios": [
            ("17001", "Manizales"),
            ("17174", "Chinchiná"),
            ("17486", "Neira"),
            ("17524", "Palestina"),
            ("17873", "Villamaría")
        ]
    },
    {
        "codigo": "AO",
        "nombre": "Alto Occidente",
        "municipios": [
            ("17272", "Filadelfia"),
            ("17388", "La Merced"),
            ("17442", "Marmato"),
            ("17614", "Riosucio"),
            ("17777", "Supía")
        ]
    },
    {
        "codigo": "BO",
        "nombre": "Bajo Occidente",
        "municipios": [
            ("17042", "Anserma"),
            ("17088", "Belalcázar"),
            ("17616", "Risaralda"),
            ("17665", "San José"),
            ("17877", "Viterbo")
        ]
    },
    {
        "codigo": "MC",
        "nombre": "Magdalena Caldense",
        "municipios": [
            ("17380", "La Dorada"),
            ("17495", "Norcasia"),
            ("17662", "Samaná"),
            ("17867", "Victoria")
        ]
    },
    {
        "codigo": "OR",
        "nombre": "Alto Oriente",
        "municipios": [
            ("17433", "Manzanares"),
            ("17444", "Marquetalia"),
            ("17446", "Marulanda"),
            ("17541", "Pensilvania")
        ]
    },
    {
        "codigo": "NT",
        "nombre": "Norte",
        "municipios": [
            ("17013", "Aguadas"),
            ("17050", "Aranzazu"),
            ("17513", "Pácora"),
            ("17653", "Salamina")
        ]
    }
]

def seed_database():
    print("Creando tablas en la base de datos...", flush=True)
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)

    db: Session = SessionLocal()

    try:
        # 1. Territorio de Caldas
        print("Cargando división político-administrativa de Caldas...", flush=True)
        dep = Departamento(codigo_dane="17", nombre="Caldas")
        db.add(dep)
        db.commit()
        db.refresh(dep)

        muni_map = {}
        for sub_data in CALDAS_SUBREGIONES:
            sub = Subregion(
                departamento_id=dep.id,
                codigo=sub_data["codigo"],
                nombre=sub_data["nombre"]
            )
            db.add(sub)
            db.commit()
            db.refresh(sub)

            for cod_dane, nom_muni in sub_data["municipios"]:
                m = Municipio(
                    subregion_id=sub.id,
                    codigo_dane=cod_dane,
                    nombre=nom_muni
                )
                db.add(m)
                db.commit()
                db.refresh(m)
                muni_map[cod_dane] = m

        # 2. Instituciones y Sedes Educativas Representativas
        print("Cargando catálogo oficial de Instituciones Educativas y Sedes de Caldas...", flush=True)
        ie_data = [
            ("117001000123", "INSTITUCIÓN EDUCATIVA UNIVERSITARIO DE CALDAS", "17001", [
                ("11700100012301", "SEDE PRINCIPAL UNIVERSITARIO", "Carrera 23 # 45-12", "Urbana"),
                ("11700100012302", "SEDE PRIMARIA CARLOS CONCHA", "Calle 48 # 22-10", "Urbana")
            ]),
            ("117001000456", "INSTITUCIÓN EDUCATIVA INSTITUTO MANIZALES", "17001", [
                ("11700100045601", "SEDE PRINCIPAL INSTITUTO MANIZALES", "Calle 14 # 21-45", "Urbana"),
                ("11700100045602", "SEDE SAN VICENTE DE PAUL", "Vereda El Arenillo", "Rural")
            ]),
            ("117001000789", "INSTITUCIÓN EDUCATIVA ESCUELA NORMAL SUPERIOR DE CALDAS", "17001", [
                ("11700100078901", "SEDE CENTRAL NORMAL", "Avenida Santander # 59-00", "Urbana")
            ]),
            ("117174000111", "INSTITUCIÓN EDUCATIVA BARTOLOMÉ MITRE", "17174", [
                ("11717400011101", "SEDE PRINCIPAL MITRE", "Calle 10 # 8-24", "Urbana"),
                ("11717400011102", "SEDE MARÍA INMACULADA", "Vereda La Quiebra", "Rural")
            ]),
            ("117380000222", "INSTITUCIÓN EDUCATIVA ALFONSO LÓPEZ", "17380", [
                ("11738000022201", "SEDE PRINCIPAL ALFONSO LÓPEZ", "Carrera 4 # 12-50", "Urbana"),
                ("11738000022202", "SEDE GUAYINAS", "Vereda Las Guayinas", "Rural")
            ]),
            ("117614000333", "INSTITUCIÓN EDUCATIVA SAN JERÓNIMO", "17614", [
                ("11761400033301", "SEDE PRINCIPAL SAN JERÓNIMO", "Calle 7 # 6-30", "Urbana"),
                ("11761400033302", "SEDE RESGUARDO CAÑAMOMO LOMAPRIETA", "Resguardo Indígena Cañamomo", "Rural")
            ]),
            ("117013000444", "INSTITUCIÓN EDUCATIVA ESCUELA NORMAL SUPERIOR CLAUDINA MÚNERA", "17013", [
                ("11701300044401", "SEDE CENTRAL AGUADAS", "Carrera 4 # 7-15", "Urbana")
            ]),
            ("117541000555", "INSTITUCIÓN EDUCATIVA PENSILVANIA", "17541", [
                ("11754100055501", "SEDE INTEGRADA", "Calle 4 # 5-20", "Urbana"),
                ("11754100055502", "SEDE VEREDA LA MERCED", "Vereda La Merced Alta", "Rural")
            ]),
            ("117042000666", "INSTITUCIÓN EDUCATIVA DE OCCIDENTE", "17042", [
                ("11704200066601", "SEDE PRINCIPAL DE OCCIDENTE", "Carrera 5 # 11-40", "Urbana")
            ])
        ]

        sedes_creadas = []
        for cod_ie, nom_ie, cod_muni, sedes in ie_data:
            muni = muni_map[cod_muni]
            ie = InstitucionEducativa(
                municipio_id=muni.id,
                codigo_oficial=cod_ie,
                nombre=nom_ie,
                rector="Directivo Docente Encargado",
                estado="Activo"
            )
            db.add(ie)
            db.commit()
            db.refresh(ie)

            for s_cod, s_nom, s_dir, s_zona in sedes:
                sede = SedeEducativa(
                    institucion_id=ie.id,
                    municipio_id=muni.id,
                    codigo_dane=s_cod,
                    nombre=s_nom,
                    direccion_vereda=s_dir,
                    zona=s_zona,
                    estado="Activa"
                )
                db.add(sede)
                db.commit()
                db.refresh(sede)
                sedes_creadas.append(sede)

        # 3. Importación desde BASE DE DATOS 2026_S2_SV.xlsm
        excel_path = "BASE DE DATOS 2026_S2_SV.xlsm"
        print(f"Abriendo libro Excel '{excel_path}'...", flush=True)
        wb = openpyxl.load_workbook(excel_path, data_only=True)

        # Capitulos
        print("Importando capítulos...", flush=True)
        ws_cap = wb['Capitulos']
        cap_obj_map = {}
        for r in range(3, ws_cap.max_row + 1):
            cod = ws_cap.cell(row=r, column=1).value
            desc = ws_cap.cell(row=r, column=2).value
            cat = ws_cap.cell(row=r, column=3).value
            orden = ws_cap.cell(row=r, column=5).value
            if cod and desc:
                cod_str = str(cod).strip()
                cap = Capitulo(
                    codigo=cod_str,
                    descripcion=str(desc).strip(),
                    categoria=str(cat or "Capitulo").strip(),
                    orden=int(orden) if orden is not None and str(orden).isdigit() else r - 2
                )
                db.add(cap)
                cap_obj_map[cod_str] = cap
        db.commit()
        for k, v in cap_obj_map.items():
            db.refresh(v)

        # Insumos y Precios 2026/2025
        print("Importando catálogo de insumos y precios 2025/2026...", flush=True)
        ws_ins = wb['Insumos']
        ins_raw = []
        for r in range(4, ws_ins.max_row + 1):
            cod = ws_ins.cell(row=r, column=1).value
            desc = ws_ins.cell(row=r, column=2).value
            tipo = ws_ins.cell(row=r, column=3).value
            und = ws_ins.cell(row=r, column=4).value
            vr_iva = ws_ins.cell(row=r, column=5).value
            vr_2025 = ws_ins.cell(row=r, column=6).value
            fuente = ws_ins.cell(row=r, column=7).value
            obs = ws_ins.cell(row=r, column=8).value
            if cod and desc:
                ins_raw.append((str(cod).strip(), str(desc).strip(), str(tipo or "Insumos").strip(), str(und or "UND").strip(), vr_iva, vr_2025, fuente, obs))

        ins_id_map = {}
        for cod_str, desc_str, tipo_str, und_str, vr_iva, vr_2025, fte, obs in ins_raw:
            ins = Insumo(
                codigo=cod_str,
                descripcion=desc_str,
                tipo=tipo_str,
                unidad=und_str,
                fuente=str(fte or "BASE DE DATOS 2026_S2_SV").strip(),
                observaciones=str(obs or "").strip()
            )
            db.add(ins)
        db.commit()

        # Build in-memory map of insumo id
        all_ins = db.query(Insumo.id, Insumo.codigo).all()
        ins_id_map = {item.codigo: item.id for item in all_ins}

        # Prices
        prices_to_add = []
        for cod_str, desc_str, tipo_str, und_str, vr_iva, vr_2025, fte, obs in ins_raw:
            ins_id = ins_id_map.get(cod_str)
            if not ins_id:
                continue
            if vr_iva is not None:
                p_num = Decimal(str(vr_iva)) if isinstance(vr_iva, (int, float, Decimal)) else Decimal("0.00")
                prices_to_add.append(PrecioInsumo(
                    insumo_id=ins_id,
                    vigencia=2026,
                    precio_unitario=p_num,
                    incluye_iva=True,
                    fuente=str(fte or "Catálogo Caldas 2026")
                ))
            if vr_2025 is not None:
                p_num25 = Decimal(str(vr_2025)) if isinstance(vr_2025, (int, float, Decimal)) else Decimal("0.00")
                prices_to_add.append(PrecioInsumo(
                    insumo_id=ins_id,
                    vigencia=2025,
                    precio_unitario=p_num25,
                    incluye_iva=True,
                    fuente="Histórico Caldas 2025"
                ))

        db.bulk_save_objects(prices_to_add)
        db.commit()
        print(f"Insumos cargados: {len(ins_id_map)}, precios: {len(prices_to_add)}", flush=True)

        # Actividades
        print("Importando catálogo maestro de actividades...", flush=True)
        ws_act = wb['Actividades']
        act_raw = []
        for r in range(9, ws_act.max_row + 1):
            cod = ws_act.cell(row=r, column=1).value
            desc = ws_act.cell(row=r, column=2).value
            cat = ws_act.cell(row=r, column=3).value
            und = ws_act.cell(row=r, column=4).value
            vr_unit = ws_act.cell(row=r, column=5).value
            if cod and desc:
                prec = Decimal(str(vr_unit)) if isinstance(vr_unit, (int, float, Decimal)) else Decimal("0.00")
                act_raw.append(Actividad(
                    codigo=str(cod).strip(),
                    descripcion=str(desc).strip(),
                    categoria=str(cat or "Actividad").strip(),
                    unidad=str(und or "UND").strip(),
                    precio_referencia=prec
                ))
        db.bulk_save_objects(act_raw)
        db.commit()

        # Build in-memory map of actividad id
        all_acts = db.query(Actividad.id, Actividad.codigo, Actividad.precio_referencia).all()
        act_map = {item.codigo: (item.id, item.precio_referencia) for item in all_acts}
        print(f"Actividades cargadas: {len(act_map)}", flush=True)

        # APUs e Insumos vinculados
        print("Importando desglose de APUs y componentes (optimizado)...", flush=True)
        ws_apu = wb['A.P.U']
        apu_records_to_create = []
        components_raw = []

        current_act_id = None
        current_apu_idx = -1

        apu_parent_rows = {}
        for r in range(2, ws_apu.max_row + 1):
            cod_act = ws_apu.cell(row=r, column=1).value
            cod_ins = ws_apu.cell(row=r, column=2).value
            tipo = ws_apu.cell(row=r, column=4).value
            vr_tot = ws_apu.cell(row=r, column=9).value

            if not cod_act:
                continue

            cod_act_str = str(cod_act).strip()

            if cod_ins == '-' or tipo in ('Analisis Basico', 'Actividad', 'Capitulo'):
                if cod_act_str in act_map:
                    act_id, default_ref = act_map[cod_act_str]
                    cost_tot = Decimal(str(vr_tot)) if isinstance(vr_tot, (int, float, Decimal)) else default_ref
                    apu_parent_rows[cod_act_str] = (act_id, cost_tot)
            elif cod_ins:
                # Component row
                cod_ins_str = str(cod_ins).strip()
                rend = ws_apu.cell(row=r, column=6).value
                p_unit = ws_apu.cell(row=r, column=7).value
                parcial = ws_apu.cell(row=r, column=8).value

                rend_dec = Decimal(str(rend)) if isinstance(rend, (int, float, Decimal)) else Decimal("1.000000")
                p_unit_dec = Decimal(str(p_unit)) if isinstance(p_unit, (int, float, Decimal)) else Decimal("0.0000")
                parcial_dec = Decimal(str(parcial)) if isinstance(parcial, (int, float, Decimal)) else Decimal("0.0000")

                components_raw.append((cod_act_str, cod_ins_str, str(tipo or "Insumos").strip(), rend_dec, p_unit_dec, parcial_dec))

        # Insert parent APUs
        apus_to_insert = [
            APU(actividad_id=act_id, version=1, costo_directo_total=cost_tot, estado="Activo")
            for cod, (act_id, cost_tot) in apu_parent_rows.items()
        ]
        db.bulk_save_objects(apus_to_insert)
        db.commit()

        # Fetch APU IDs mapped by activity id
        all_apus = db.query(APU.id, APU.actividad_id).all()
        apu_id_by_act_id = {item.actividad_id: item.id for item in all_apus}

        # Build components
        components_to_insert = []
        for cod_act_str, cod_ins_str, tipo_str, rend_dec, p_unit_dec, parcial_dec in components_raw:
            if cod_act_str not in act_map:
                continue
            act_id, _ = act_map[cod_act_str]
            apu_id = apu_id_by_act_id.get(act_id)
            if not apu_id:
                continue

            ins_id = ins_id_map.get(cod_ins_str)
            if ins_id:
                components_to_insert.append(APUComponente(
                    apu_id=apu_id,
                    insumo_id=ins_id,
                    tipo_componente=tipo_str,
                    rendimiento=rend_dec,
                    precio_aplicado=p_unit_dec,
                    costo_parcial=parcial_dec
                ))

        db.bulk_save_objects(components_to_insert)
        db.commit()
        print(f"APUs padres creados: {len(apu_parent_rows)}, componentes creados: {len(components_to_insert)}", flush=True)

        # 4. Parámetros de Costos Indirectos (AIU)
        print("Cargando estructura estándar de AIU...", flush=True)
        aiu_defaults = [
            ("1. GASTOS GENERALES", "Ingeniero Director", "Mes", 5, 1, 6700000, 0.6556),
            ("1. GASTOS GENERALES", "Ingeniero Residente", "Mes", 5, 1, 4500000, 0.6556),
            ("1. GASTOS GENERALES", "Maestro de Obra", "Mes", 5, 1, 3600000, 0.6556),
            ("1. GASTOS GENERALES", "Residente de Salud Ocupacional", "Mes", 5, 1, 3600000, 0.6556),
            ("1. GASTOS GENERALES", "Campamento y Oficina de Obra", "Mes", 5, 1, 800000, 0.0),
            ("1. GASTOS GENERALES", "Sistema de Gestión Ambiental y Ensayos de Calidad", "Mes", 5, 1, 3000000, 0.0),
            ("2. GASTOS INDIRECTOS", "Póliza de Cumplimiento y Responsabilidad Civil", "Global", 1, 1, 4500000, 0.0),
            ("2. GASTOS INDIRECTOS", "Vallas y Señalización Institucional", "M2", 2, 1, 300000, 0.0),
            ("2. GASTOS INDIRECTOS", "Impuestos y Estampillas Departamentales", "Global", 1, 1, 12500000, 0.0)
        ]

        for cap_ind, conc, und, cant, ded, vr_u, fp in aiu_defaults:
            tot = Decimal(str(cant * ded * vr_u * (1 + fp)))
            db.add(AnalisisIndirecto(
                presupuesto_id=None,
                capitulo_indirecto=cap_ind,
                concepto=conc,
                unidad=und,
                cantidad=Decimal(str(cant)),
                dedicacion=Decimal(str(ded)),
                valor_unitario=Decimal(str(vr_u)),
                factor_prestacional_aiu=Decimal(str(fp)),
                valor_total=tot
            ))

        # 5. Parámetros Prestacionales Legales
        print("Cargando factores prestacionales...", flush=True)
        prestacionales = [
            ("Cesantías", "Prestaciones Sociales", "Alto", "1-2 SMMLV", 0.0833),
            ("Intereses de Cesantías", "Prestaciones Sociales", "Alto", "1-2 SMMLV", 0.0100),
            ("Prima de Servicios", "Prestaciones Sociales", "Alto", "1-2 SMMLV", 0.0833),
            ("Vacaciones", "Prestaciones Sociales", "Alto", "1-2 SMMLV", 0.0417),
            ("Dotaciones", "Prestaciones Sociales", "Alto", "1-2 SMMLV", 0.0600),
            ("Pensión Patronal", "Seguridad Social", "Alto", "1-2 SMMLV", 0.1200),
            ("ARL (Nivel V Obra)", "Seguridad Social", "Alto", "1-2 SMMLV", 0.0696),
            ("Caja de Compensación Familiar", "Parafiscales", "Alto", "1-2 SMMLV", 0.0400),
            ("FIC (Fondo Industria Construcción)", "Parafiscales", "Alto", "1-2 SMMLV", 0.0250)
        ]
        for conc, cat, riesg, rang, val in prestacionales:
            db.add(ParametroPrestacional(
                concepto=conc,
                categoria=cat,
                riesgo=riesg,
                rango_salarial=rang,
                valor=Decimal(str(val)),
                vigencia=2026,
                fuente="Normativa Laboral Colombiana 2026"
            ))

        # 6. Pólizas Contractuales
        print("Cargando rangos de pólizas...", flush=True)
        polizas_seed = [
            ("Responsabilidad Civil Extracontractual (RCE)", 0, 2135250000, 200, 0.001413, 402423),
            ("Responsabilidad Civil Extracontractual (RCE)", 2135250000, 3558750000, 300, 0.000902, 385319),
            ("Cumplimiento del Contrato", 0, 5000000000, 100, 0.001500, 350000),
            ("Buen Manejo del Anticipo", 0, 5000000000, 100, 0.001200, 280000),
            ("Estabilidad de la Obra", 0, 5000000000, 100, 0.001800, 420000)
        ]
        for tipo_p, l_inf, l_sup, smm, tf, tf_f in polizas_seed:
            db.add(PolizaRango(
                tipo_poliza=tipo_p,
                limite_inferior=Decimal(str(l_inf)),
                limite_superior=Decimal(str(l_sup)) if l_sup else None,
                cobertura_smmlv=Decimal(str(smm)),
                tarifa_valor=Decimal(str(tf)),
                tarifa_fija=Decimal(str(tf_f)),
                vigencia=2026,
                fuente="Resolución Pólizas Obras Caldas 2026"
            ))

        # 7. Usuarios del Sistema
        usuarios_seed = [
            ("Ing. Andrés Gómez (Secretaría de Educación)", "andres.gomez@caldas.gov.co", "Administrador"),
            ("Arq. Marcela Cárdenas (Formulación Obras)", "marcela.cardenas@caldas.gov.co", "Formulador"),
            ("Ing. Carlos Echeverry (Control Costos)", "carlos.echeverry@caldas.gov.co", "Administrador de Costos"),
            ("Dra. Beatriz Restrepo (Revisión Técnica)", "beatriz.restrepo@caldas.gov.co", "Revisor"),
            ("Dr. Luis Fernando Patiño (Secretario Aprobador)", "luis.patino@caldas.gov.co", "Aprobador"),
            ("Dra. Claudia Montoya (Control Interno)", "claudia.montoya@caldas.gov.co", "Auditor"),
            ("Veeduría Ciudadana Caldas", "veeduria.educativa@caldas.org", "Consulta")
        ]
        for nom, em, rol in usuarios_seed:
            db.add(Usuario(nombre=nom, email=em, rol=rol, activo=True))
        db.commit()

        # 8. Proyectos y Presupuestos Semilla
        print("Creando proyectos demostrativos y presupuestos modelo...", flush=True)
        # Proyecto 1: Manizales I.E. Universitario de Caldas
        sede_uni = sedes_creadas[0] # Sede Principal Universitario
        p1 = ProyectoObra(
            codigo="CALDAS-EDU-2026-001",
            nombre="Construcción Bloque de 6 Aulas y Laboratorio STEM — I.E. Universitario de Caldas",
            tipo_intervencion="Construcción",
            objeto="Construcción de infraestructura educativa en 2 niveles con estructura sismorresistente, baterías sanitarias y redes técnicas.",
            estado="Aprobado",
            responsable="Arq. Marcela Cárdenas",
            fuente_financiacion="Sistema General de Regalías (SGR) & Recursos Propios Caldas",
            fecha_inicio=datetime.date(2026, 2, 1),
            fecha_fin=datetime.date(2026, 8, 31)
        )
        db.add(p1)
        db.commit()
        db.refresh(p1)

        db.add(ProyectoSede(
            proyecto_id=p1.id,
            sede_id=sede_uni.id,
            alcance_sede="Construcción bloque aulas sede principal Manizales"
        ))
        db.commit()

        # Presupuesto para Proyecto 1
        pres1 = Presupuesto(
            proyecto_id=p1.id,
            version=1,
            vigencia=2026,
            estado="Borrador",
            porcentaje_aiu=Decimal("0.25")
        )
        db.add(pres1)
        db.commit()
        db.refresh(pres1)

        cap_prelim = db.query(Capitulo).filter(Capitulo.codigo.in_(["001", "1"])).first() or db.query(Capitulo).first()
        cap_tierra = db.query(Capitulo).filter(Capitulo.codigo.in_(["002", "2"])).first() or cap_prelim
        cap_concreto = db.query(Capitulo).filter(Capitulo.codigo.in_(["005", "5"])).first() or cap_prelim

        sample_acts = db.query(Actividad).filter(Actividad.precio_referencia > 0).limit(12).all()
        for idx, act in enumerate(sample_acts, start=1):
            cap_assigned = cap_prelim if idx <= 4 else (cap_tierra if idx <= 8 else cap_concreto)
            db.add(PresupuestoItem(
                presupuesto_id=pres1.id,
                capitulo_id=cap_assigned.id,
                actividad_id=act.id,
                item_num=f"{cap_assigned.codigo}.{idx}",
                cantidad=Decimal(str(25 + idx * 10)),
                precio_unitario=act.precio_referencia,
                subtotal=Decimal("0.00"),
                orden=idx,
                observaciones=f"Partida contractual para obra {p1.codigo}"
            ))
        db.commit()

        # Recalcular y APROBAR (Congelar) Presupuesto 1
        recalculate_budget(pres1.id, db)
        freeze_budget_snapshot(pres1.id, "Dr. Luis Fernando Patiño (Secretario Aprobador)", db)

        # Proyecto 2: La Dorada I.E. Alfonso López (En Revisión)
        sede_dorada = next((s for s in sedes_creadas if "ALFONSO" in s.nombre), sedes_creadas[1])
        p2 = ProyectoObra(
            codigo="CALDAS-EDU-2026-002",
            nombre="Adecuación Integral y Baterías Sanitarias — I.E. Alfonso López",
            tipo_intervencion="Adecuación",
            objeto="Mejoramiento hidrosanitario, reposición de pisos y adecuación de baterías para inclusión de personas con movilidad reducida.",
            estado="En Revisión",
            responsable="Ing. Carlos Echeverry",
            fuente_financiacion="Sistema General de Participaciones (SGP) Educación",
            fecha_inicio=datetime.date(2026, 3, 1),
            fecha_fin=datetime.date(2026, 7, 15)
        )
        db.add(p2)
        db.commit()
        db.refresh(p2)

        db.add(ProyectoSede(
            proyecto_id=p2.id,
            sede_id=sede_dorada.id,
            alcance_sede="Adecuación sanitaria integral sede urbana La Dorada"
        ))
        db.commit()

        pres2 = Presupuesto(
            proyecto_id=p2.id,
            version=1,
            vigencia=2026,
            estado="En Revisión",
            porcentaje_aiu=Decimal("0.25")
        )
        db.add(pres2)
        db.commit()
        db.refresh(pres2)

        for idx, act in enumerate(sample_acts[:8], start=1):
            db.add(PresupuestoItem(
                presupuesto_id=pres2.id,
                capitulo_id=cap_prelim.id,
                actividad_id=act.id,
                item_num=f"1.{idx}",
                cantidad=Decimal(str(15 + idx * 5)),
                precio_unitario=act.precio_referencia,
                subtotal=Decimal("0.00"),
                orden=idx,
                observaciones="Mantenimiento preventivo y correctivo"
            ))
        db.commit()
        recalculate_budget(pres2.id, db)

        # Proyecto 3: Riosucio Multi-Sede (Borrador)
        sede_rio1 = next((s for s in sedes_creadas if "JERÓNIMO" in s.nombre or "JERONIMO" in s.nombre), sedes_creadas[2])
        sede_rio2 = next((s for s in sedes_creadas if "CAÑAMOMO" in s.nombre or "CANAMOMO" in s.nombre), sedes_creadas[3])
        p3 = ProyectoObra(
            codigo="CALDAS-EDU-2026-003",
            nombre="Mantenimiento de Cubiertas y Placas Polideportivas — I.E. San Jerónimo",
            tipo_intervencion="Mantenimiento",
            objeto="Impermeabilización de cubiertas, reposición de caballetes y mantenimiento de escenarios deportivos escolares.",
            estado="Borrador",
            responsable="Arq. Marcela Cárdenas",
            fuente_financiacion="Recursos Propios Municipio de Riosucio y Caldas",
            fecha_inicio=datetime.date(2026, 4, 1),
            fecha_fin=datetime.date(2026, 9, 30)
        )
        db.add(p3)
        db.commit()
        db.refresh(p3)

        db.add(ProyectoSede(proyecto_id=p3.id, sede_id=sede_rio1.id, alcance_sede="Sede Urbana Riosucio"))
        db.add(ProyectoSede(proyecto_id=p3.id, sede_id=sede_rio2.id, alcance_sede="Sede Rural Resguardo"))
        db.commit()

        pres3 = Presupuesto(
            proyecto_id=p3.id,
            version=1,
            vigencia=2026,
            estado="Borrador",
            porcentaje_aiu=Decimal("0.25")
        )
        db.add(pres3)
        db.commit()
        db.refresh(pres3)

        for idx, act in enumerate(sample_acts[4:10], start=1):
            db.add(PresupuestoItem(
                presupuesto_id=pres3.id,
                capitulo_id=cap_prelim.id,
                actividad_id=act.id,
                item_num=f"1.{idx}",
                cantidad=Decimal(str(30 + idx * 8)),
                precio_unitario=act.precio_referencia,
                subtotal=Decimal("0.00"),
                orden=idx
            ))
        db.commit()
        recalculate_budget(pres3.id, db)

        print("\n¡Sembrado de base de datos completado exitosamente!", flush=True)
        print(f"- Subregiones: {len(CALDAS_SUBREGIONES)}", flush=True)
        print(f"- Municipios: {len(muni_map)}", flush=True)
        print(f"- Instituciones Educativas: {len(ie_data)}", flush=True)
        print(f"- Sedes Educativas: {len(sedes_creadas)}", flush=True)
        print(f"- Capítulos importados: {len(cap_obj_map)}", flush=True)
        print(f"- Insumos importados: {len(ins_id_map)}", flush=True)
        print(f"- Actividades importadas: {len(act_map)}", flush=True)
        print(f"- Proyectos configurados: 3 (1 Aprobado/Congelado, 1 En Revisión, 1 Borrador)", flush=True)

    except Exception as e:
        db.rollback()
        print(f"Error sembrando datos: {e}", flush=True)
        raise e
    finally:
        db.close()

if __name__ == "__main__":
    seed_database()
