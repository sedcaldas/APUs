import io
from decimal import Decimal
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from sqlalchemy.orm import Session
from backend.app.models import Presupuesto

def export_budget_excel(budget_id: int, db: Session) -> io.BytesIO:
    """Exporta el presupuesto de obra pública detallado a formato Excel oficial."""
    presupuesto = db.query(Presupuesto).filter(Presupuesto.id == budget_id).first()
    if not presupuesto:
        raise ValueError(f"Presupuesto {budget_id} no encontrado")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Presupuesto_Detallado"

    # Color Palette: Gobernación de Caldas
    primary_color = "1B365D" # Deep Navy
    accent_color = "00805A"  # Caldas Emerald
    light_gray = "F4F6F9"
    border_gray = "D1D5DB"

    fill_header = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    font_header = Font(color="FFFFFF", bold=True, size=11)
    fill_capitulo = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    font_capitulo = Font(color="1E293B", bold=True, size=10)

    font_title = Font(color=primary_color, bold=True, size=14)
    font_sub = Font(color="4B5563", size=10)
    font_bold = Font(bold=True)

    border_thin = Border(
        left=Side(style='thin', color=border_gray),
        right=Side(style='thin', color=border_gray),
        top=Side(style='thin', color=border_gray),
        bottom=Side(style='thin', color=border_gray)
    )

    # 1. Institutional Header
    ws.merge_cells("A1:G1")
    ws["A1"] = "GOBERNACIÓN DE CALDAS — SECRETARÍA DE EDUCACIÓN"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:G2")
    ws["A2"] = f"PRESUPUESTO OFICIAL DE OBRA PÚBLICA — INFRAESTRUCTURA EDUCATIVA"
    ws["A2"].font = Font(color="1E293B", bold=True, size=12)
    ws["A2"].alignment = Alignment(horizontal="center")

    # Project metadata
    proy = presupuesto.proyecto
    sedes_nombres = ", ".join([ps.sede.nombre for ps in proy.sedes_asociadas]) if proy.sedes_asociadas else "N/A"
    
    metadata = [
        ("Proyecto:", proy.nombre, "Código:", proy.codigo),
        ("Tipo Intervención:", proy.tipo_intervencion, "Vigencia:", str(presupuesto.vigencia)),
        ("Sede(s) Educativa(s):", sedes_nombres, "Versión Presupuesto:", f"v{presupuesto.version}"),
        ("Estado:", presupuesto.estado.upper(), "Moneda:", presupuesto.moneda)
    ]

    r_idx = 4
    for row in metadata:
        ws.cell(row=r_idx, column=1, value=row[0]).font = font_bold
        ws.cell(row=r_idx, column=2, value=row[1])
        ws.cell(row=r_idx, column=5, value=row[2]).font = font_bold
        ws.cell(row=r_idx, column=6, value=row[3])
        r_idx += 1

    r_idx += 1

    # Table Header
    headers = ["ÍTEM", "CÓDIGO", "DESCRIPCIÓN DE LA ACTIVIDAD", "UNIDAD", "CANTIDAD", "VALOR UNITARIO", "VALOR TOTAL"]
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=r_idx, column=col_idx, value=h)
        c.fill = fill_header
        c.font = font_header
        c.alignment = Alignment(horizontal="center" if col_idx != 3 else "left", vertical="center")
        c.border = border_thin

    # Group items by Chapter
    items = sorted(presupuesto.items, key=lambda x: (x.capitulo.orden, x.orden))
    current_cap = None
    item_counter = 1

    for it in items:
        if it.capitulo_id != current_cap:
            current_cap = it.capitulo_id
            r_idx += 1
            ws.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=7)
            cap_cell = ws.cell(row=r_idx, column=1, value=f"CAPÍTULO {it.capitulo.codigo}: {it.capitulo.descripcion}")
            cap_cell.fill = fill_capitulo
            cap_cell.font = font_capitulo
            cap_cell.border = border_thin

        r_idx += 1
        ws.cell(row=r_idx, column=1, value=it.item_num or str(item_counter)).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=2, value=it.actividad.codigo).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=3, value=it.actividad.descripcion)
        ws.cell(row=r_idx, column=4, value=it.actividad.unidad or "UND").alignment = Alignment(horizontal="center")
        
        c_cant = ws.cell(row=r_idx, column=5, value=float(it.cantidad))
        c_cant.number_format = "#,##0.00"
        
        c_prec = ws.cell(row=r_idx, column=6, value=float(it.precio_unitario))
        c_prec.number_format = "$#,##0.00"

        c_sub = ws.cell(row=r_idx, column=7, value=float(it.subtotal))
        c_sub.number_format = "$#,##0.00"

        for c_i in range(1, 8):
            ws.cell(row=r_idx, column=c_i).border = border_thin
        item_counter += 1

    r_idx += 2

    # Financial Summary
    summary_data = [
        ("COSTO DIRECTO TOTAL", float(presupuesto.costo_directo)),
        (f"COSTOS INDIRECTOS (AIU {float(presupuesto.porcentaje_aiu * 100):.1f}%)", float(presupuesto.costo_indirecto)),
        ("VALOR TOTAL PRESUPUESTO", float(presupuesto.valor_total))
    ]

    for label, val in summary_data:
        ws.merge_cells(start_row=r_idx, start_column=4, end_row=r_idx, end_column=6)
        c_lbl = ws.cell(row=r_idx, column=4, value=label)
        c_lbl.font = font_bold
        c_lbl.alignment = Alignment(horizontal="right")
        
        c_val = ws.cell(row=r_idx, column=7, value=val)
        c_val.font = font_bold
        c_val.number_format = "$#,##0.00"
        if "VALOR TOTAL" in label:
            c_lbl.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
            c_val.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        r_idx += 1

    # Column Widths
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 45
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['F'].width = 18
    ws.column_dimensions['G'].width = 20

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def export_budget_pdf(budget_id: int, db: Session) -> io.BytesIO:
    """Genera informe en PDF para impresión y control de interventoría."""
    presupuesto = db.query(Presupuesto).filter(Presupuesto.id == budget_id).first()
    if not presupuesto:
        raise ValueError(f"Presupuesto {budget_id} no encontrado")

    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    story = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=14,
        textColor=colors.HexColor('#1B365D'),
        alignment=1, # Centered
        spaceAfter=6
    )
    sub_style = ParagraphStyle(
        'SubTitle',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.HexColor('#4B5563'),
        alignment=1,
        spaceAfter=14
    )

    story.append(Paragraph("GOBERNACIÓN DE CALDAS — SECRETARÍA DE EDUCACIÓN", title_style))
    story.append(Paragraph(f"RESUMEN EJECUTIVO DE PRESUPUESTO — PROYECTO: {presupuesto.proyecto.nombre}", sub_style))

    # Info table
    info_data = [
        ["Código Proyecto:", presupuesto.proyecto.codigo, "Vigencia:", str(presupuesto.vigencia)],
        ["Intervención:", presupuesto.proyecto.tipo_intervencion, "Versión:", f"v{presupuesto.version}"],
        ["Estado:", presupuesto.estado.upper(), "Fecha:", str(presupuesto.updated_at.date() if presupuesto.updated_at else "")]
    ]
    t_info = Table(info_data, colWidths=[110, 160, 90, 180])
    t_info.setStyle(TableStyle([
        ('FONTNAME', (0,0), (-1,-1), 'Helvetica'),
        ('FONTSIZE', (0,0), (-1,-1), 8),
        ('FONTNAME', (0,0), (0,-1), 'Helvetica-Bold'),
        ('FONTNAME', (2,0), (2,-1), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ('TOPPADDING', (0,0), (-1,-1), 3),
    ]))
    story.append(t_info)
    story.append(Spacer(1, 12))

    # Items Table
    table_data = [["Ítem", "Código", "Descripción", "Und", "Cant.", "Vr. Unit.", "Subtotal"]]
    for idx, it in enumerate(presupuesto.items[:60], start=1):
        desc = it.actividad.descripcion[:38] + ("..." if len(it.actividad.descripcion) > 38 else "")
        table_data.append([
            str(it.item_num or idx),
            it.actividad.codigo,
            desc,
            it.actividad.unidad or "UND",
            f"{float(it.cantidad):,.2f}",
            f"${float(it.precio_unitario):,.0f}",
            f"${float(it.subtotal):,.0f}"
        ])

    t_items = Table(table_data, colWidths=[30, 50, 220, 35, 55, 70, 80])
    t_items.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1B365D')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 7),
        ('ALIGN', (3,1), (3,-1), 'CENTER'),
        ('ALIGN', (4,1), (-1,-1), 'RIGHT'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#D1D5DB')),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ('TOPPADDING', (0,0), (-1,-1), 3),
    ]))
    story.append(t_items)
    story.append(Spacer(1, 14))

    # Totals Box
    totals_data = [
        ["COSTO DIRECTO:", f"${float(presupuesto.costo_directo):,.2f}"],
        [f"COSTOS INDIRECTOS (AIU {float(presupuesto.porcentaje_aiu * 100):.1f}%):", f"${float(presupuesto.costo_indirecto):,.2f}"],
        ["VALOR TOTAL PRESUPUESTO:", f"${float(presupuesto.valor_total):,.2f}"]
    ]
    t_tot = Table(totals_data, colWidths=[380, 160])
    t_tot.setStyle(TableStyle([
        ('FONTNAME', (0,0), (-1,-1), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 9),
        ('ALIGN', (0,0), (-1,-1), 'RIGHT'),
        ('BACKGROUND', (0,2), (-1,2), colors.HexColor('#FEF3C7')),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('TOPPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t_tot)

    doc.build(story)
    output.seek(0)
    return output
